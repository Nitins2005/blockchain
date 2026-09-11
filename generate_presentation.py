#!/usr/bin/env python3
"""
CryptoShield AI — Professional PowerPoint Presentation Generator
Generates a 13-slide defense presentation with dark cybersecurity theme.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
import os

# ─── Theme Colors ────────────────────────────────────────────────────────────
BG_DARK       = RGBColor(0x0B, 0x0F, 0x19)
BG_CARD       = RGBColor(0x11, 0x18, 0x27)
BG_CARD_ALT   = RGBColor(0x1E, 0x29, 0x3B)
CYAN          = RGBColor(0x00, 0xF2, 0xFE)
CYAN_DIM      = RGBColor(0x38, 0xBD, 0xF8)
GREEN         = RGBColor(0x10, 0xB9, 0x81)
GREEN_LIGHT   = RGBColor(0x34, 0xD3, 0x99)
RED           = RGBColor(0xEF, 0x44, 0x44)
ORANGE        = RGBColor(0xF5, 0x9E, 0x0B)
WHITE         = RGBColor(0xF1, 0xF5, 0xF9)
GRAY          = RGBColor(0x94, 0xA3, 0xB8)
GRAY_LIGHT    = RGBColor(0xCB, 0xD5, 0xE1)
GRAY_DARK     = RGBColor(0x64, 0x74, 0x8B)
PURPLE        = RGBColor(0xA7, 0x8B, 0xFA)
BLUE          = RGBColor(0x60, 0xA5, 0xFA)
PINK          = RGBColor(0xF4, 0x72, 0xB6)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

BLANK_LAYOUT = prs.slide_layouts[6]  # blank layout

# ─── Helper Functions ────────────────────────────────────────────────────────

def set_slide_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, fill_color=None, line_color=None, line_width=Pt(0)):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color or BG_CARD
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()
    # Reduce corner rounding
    shape.adjustments[0] = 0.05
    return shape

def add_rect(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color or BG_CARD
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, text, left, top, width, height,
             font_size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT,
             font_name="Calibri", line_spacing=1.2):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    if line_spacing:
        p.line_spacing = Pt(font_size * line_spacing)
    return txBox

def add_multiline(slide, lines, left, top, width, height,
                  font_size=16, color=WHITE, align=PP_ALIGN.LEFT,
                  font_name="Calibri", bold=False, line_spacing=1.5):
    """lines: list of (text, color, bold, font_size) or just strings"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line_data in enumerate(lines):
        if isinstance(line_data, str):
            txt, col, bld, fs = line_data, color, bold, font_size
        elif len(line_data) == 2:
            txt, col = line_data
            bld, fs = bold, font_size
        elif len(line_data) == 3:
            txt, col, bld = line_data
            fs = font_size
        else:
            txt, col, bld, fs = line_data
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = txt
        p.font.size = Pt(fs)
        p.font.bold = bld
        p.font.color.rgb = col
        p.font.name = font_name
        p.alignment = align
        p.space_after = Pt(2)
        p.line_spacing = Pt(fs * line_spacing)
    return txBox

def add_accent_line(slide, left, top, width, color=CYAN):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_slide_number(slide, num, total=13):
    add_text(slide, f"{num} / {total}", Inches(12.2), Inches(7.05),
             Inches(1), Inches(0.4), font_size=10, color=GRAY_DARK,
             align=PP_ALIGN.RIGHT)

def add_section_header(slide, number, title, subtitle=""):
    add_text(slide, f"0{number}" if number < 10 else str(number),
             Inches(0.7), Inches(0.4), Inches(1), Inches(0.5),
             font_size=14, bold=True, color=CYAN, font_name="Consolas")
    add_text(slide, title, Inches(0.7), Inches(0.75), Inches(10), Inches(0.6),
             font_size=32, bold=True, color=WHITE, font_name="Calibri")
    add_accent_line(slide, Inches(0.7), Inches(1.35), Inches(2.5), CYAN)
    if subtitle:
        add_text(slide, subtitle, Inches(0.7), Inches(1.5), Inches(10), Inches(0.4),
                 font_size=14, color=GRAY, font_name="Calibri")

def add_footer(slide):
    add_text(slide, "CryptoShield AI  |  Nitin S.  |  Dept. of Computer Science & Engineering",
             Inches(0.7), Inches(7.05), Inches(6), Inches(0.4),
             font_size=9, color=GRAY_DARK, font_name="Calibri")
    # Thin line
    add_rect(slide, Inches(0), Inches(6.95), SLIDE_W, Pt(1), fill_color=RGBColor(0x1E, 0x29, 0x3B))


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)

# Decorative top gradient bar
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=CYAN)

# Shield icon placeholder
shield = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(5.9), Inches(0.9), Inches(1.5), Inches(1.5))
shield.fill.solid()
shield.fill.fore_color.rgb = RGBColor(0x00, 0x3D, 0x45)
shield.line.color.rgb = CYAN
shield.line.width = Pt(2)
add_text(slide, "AI", Inches(5.9), Inches(1.2), Inches(1.5), Inches(0.8),
         font_size=28, bold=True, color=CYAN, align=PP_ALIGN.CENTER, font_name="Consolas")

add_text(slide, "CryptoShield AI", Inches(2), Inches(2.7), Inches(9.3), Inches(0.8),
         font_size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font_name="Calibri")

add_text(slide, "A Temporal Explainable Multi-Chain Graph Neural Network\nfor Cryptocurrency Fraud Detection & Wallet Attribution",
         Inches(2), Inches(3.55), Inches(9.3), Inches(0.9),
         font_size=18, color=GRAY_LIGHT, align=PP_ALIGN.CENTER)

add_accent_line(slide, Inches(5.2), Inches(4.6), Inches(3), CYAN)

add_text(slide, "Presenter: Nitin S.", Inches(2), Inches(4.9), Inches(9.3), Inches(0.4),
         font_size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, "Department of Computer Science & Engineering", Inches(2), Inches(5.3), Inches(9.3), Inches(0.3),
         font_size=13, color=GRAY, align=PP_ALIGN.CENTER)

# Metric cards
metrics = [
    ("94.2%", "Accuracy", GREEN),
    ("0.963", "AUC-ROC", CYAN),
    ("11.6ms", "Latency", PURPLE),
    ("5 Chains", "Multi-Chain", ORANGE),
]
card_w = Inches(2.1)
card_h = Inches(1.0)
start_x = Inches(1.8)
gap = Inches(0.3)
for i, (val, label, col) in enumerate(metrics):
    x = start_x + i * (card_w + gap)
    card = add_shape(slide, x, Inches(6.0), card_w, card_h, fill_color=BG_CARD, line_color=col, line_width=Pt(1.5))
    add_text(slide, val, x, Inches(6.05), card_w, Inches(0.5),
             font_size=22, bold=True, color=col, align=PP_ALIGN.CENTER, font_name="Consolas")
    add_text(slide, label, x, Inches(6.5), card_w, Inches(0.35),
             font_size=11, color=GRAY, align=PP_ALIGN.CENTER)

add_slide_number(slide, 1)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — PROBLEM STATEMENT & MOTIVATION
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 1, "Problem Statement & Motivation", "Why current fraud detection fails")

# Left card — The Problem
card = add_shape(slide, Inches(0.7), Inches(2.0), Inches(5.8), Inches(4.8), fill_color=BG_CARD)
add_text(slide, "The Threat Landscape", Inches(1.0), Inches(2.15), Inches(5.2), Inches(0.4),
         font_size=18, bold=True, color=RED)

problems = [
    ("$10B+", "Annual financial extortion from ransomware, mixers, phishing, DeFi exploits"),
    ("Cross-Chain", "Fraudsters jump between BTC, ETH, BNB, Polygon, Tron to evade tracking"),
    ("Evolving", "Peeling chains, Tornado Cash mixers, flash loan attacks bypass static rules"),
]
y = Inches(2.7)
for metric, desc in problems:
    # Metric badge
    badge = add_shape(slide, Inches(1.0), y, Inches(1.2), Inches(0.4), fill_color=RGBColor(0x7F, 0x1D, 0x1D))
    add_text(slide, metric, Inches(1.0), y + Pt(2), Inches(1.2), Inches(0.35),
             font_size=13, bold=True, color=RED, align=PP_ALIGN.CENTER, font_name="Consolas")
    add_text(slide, desc, Inches(2.4), y, Inches(3.8), Inches(0.5),
             font_size=13, color=GRAY_LIGHT)
    y += Inches(0.65)

# Right card — Legacy Flaws
card = add_shape(slide, Inches(6.8), Inches(2.0), Inches(5.8), Inches(4.8), fill_color=BG_CARD)
add_text(slide, "Legacy Approach Flaws", Inches(7.1), Inches(2.15), Inches(5.2), Inches(0.4),
         font_size=18, bold=True, color=ORANGE)

flaws = [
    ("Static Rules / Blacklists", "Easily bypassed by creating new addresses;\nno adaptability to novel patterns", RED),
    ("Tabular ML (XGBoost)", "Ignores graph topology & neighbor risk;\neach wallet analyzed in isolation", ORANGE),
    ("Black-Box Deep Models", "Lack legal & forensic explainability;\nregulators cannot audit decisions", PINK),
]
y = Inches(2.7)
for title, desc, col in flaws:
    add_rect(slide, Inches(7.1), y, Pt(4), Inches(0.8), fill_color=col)
    add_text(slide, title, Inches(7.4), y, Inches(5.0), Inches(0.3),
             font_size=14, bold=True, color=col)
    add_text(slide, desc, Inches(7.4), y + Inches(0.32), Inches(5.0), Inches(0.5),
             font_size=12, color=GRAY_LIGHT)
    y += Inches(1.05)

add_footer(slide)
add_slide_number(slide, 2)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — KEY CONTRIBUTIONS & INNOVATIONS
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 2, "Key Contributions & Innovations", "Four novel contributions to blockchain forensics")

contributions = [
    ("Multi-Chain\nGraph Synthesis", "Ingests dynamic transaction streams across 5 major blockchains (BTC, ETH, BNB, MATIC, TRX) using PostgreSQL + Neo4j hybrid storage.", CYAN, "01"),
    ("Dual-Head\nT-EGNN Architecture", "Jointly outputs Fraud Risk Scores (0-100%) and 8-Class Wallet Attribution (Exchange, Mixer, DeFi, Scam, etc.) from a unified GNN backbone.", GREEN, "02"),
    ("Embedded\nXAI Module", "Computes normalized feature attributions and sub-graph explanations in-line, without post-hoc inference delays.", PURPLE, "03"),
    ("Real-Time\nPerformance", "Achieves 11.6ms query latency, meeting sub-50ms exchange compliance rules for production deployment.", ORANGE, "04"),
]

card_w = Inches(2.8)
card_h = Inches(4.3)
start_x = Inches(0.7)
gap = Inches(0.25)

for i, (title, desc, col, num) in enumerate(contributions):
    x = start_x + i * (card_w + gap)
    y = Inches(2.0)
    card = add_shape(slide, x, y, card_w, card_h, fill_color=BG_CARD, line_color=col, line_width=Pt(1.5))
    # Number badge
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), y + Inches(0.15), Inches(0.5), Inches(0.5))
    badge.fill.solid()
    badge.fill.fore_color.rgb = col
    badge.line.fill.background()
    add_text(slide, num, x + Inches(0.15), y + Inches(0.18), Inches(0.5), Inches(0.45),
             font_size=16, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER, font_name="Consolas")
    # Title
    add_text(slide, title, x + Inches(0.15), y + Inches(0.8), card_w - Inches(0.3), Inches(0.8),
             font_size=16, bold=True, color=col, font_name="Calibri", line_spacing=1.3)
    # Description
    add_text(slide, desc, x + Inches(0.15), y + Inches(1.7), card_w - Inches(0.3), Inches(2.4),
             font_size=12, color=GRAY_LIGHT, line_spacing=1.4)

add_footer(slide)
add_slide_number(slide, 3)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — SYSTEM ARCHITECTURE
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 3, "System Architecture & Pipeline", "End-to-end multi-chain fraud detection system")

layers = [
    ("Multi-Chain Ingestion", "BTC, ETH, BNB, MATIC, TRX", CYAN, "Blockchain APIs"),
    ("Hybrid Storage Layer", "PostgreSQL + Neo4j Graph DB", BLUE, "Data Persistence"),
    ("AI Inference Engine", "PyTorch T-EGNN Model", GREEN, "ML Prediction"),
    ("FastAPI Service Layer", "REST API + JWT Auth", PURPLE, "Backend API"),
    ("React Dashboard", "Cytoscape.js Interactive Canvas", ORANGE, "Frontend UI"),
]

box_w = Inches(3.5)
box_h = Inches(0.85)
start_y = Inches(2.1)
center_x = Inches(4.9)

for i, (title, desc, col, badge_text) in enumerate(layers):
    y = start_y + i * Inches(1.05)
    # Main box
    box = add_shape(slide, center_x, y, box_w, box_h, fill_color=BG_CARD, line_color=col, line_width=Pt(2))
    add_text(slide, title, center_x + Inches(0.2), y + Inches(0.08), box_w - Inches(0.4), Inches(0.35),
             font_size=15, bold=True, color=col, font_name="Calibri")
    add_text(slide, desc, center_x + Inches(0.2), y + Inches(0.43), box_w - Inches(0.4), Inches(0.3),
             font_size=11, color=GRAY_LIGHT)
    # Badge on the left
    badge = add_shape(slide, Inches(1.5), y + Inches(0.15), Inches(2.8), Inches(0.55), fill_color=BG_CARD_ALT)
    add_text(slide, badge_text, Inches(1.5), y + Inches(0.2), Inches(2.8), Inches(0.45),
             font_size=12, bold=True, color=col, align=PP_ALIGN.CENTER, font_name="Consolas")
    # Badge on the right
    if i == 0:
        right_text = "5 Chains"
    elif i == 1:
        right_text = "12 Tables"
    elif i == 2:
        right_text = "94.2% Acc"
    elif i == 3:
        right_text = "14 Routers"
    else:
        right_text = "13 Pages"
    badge_r = add_shape(slide, Inches(8.9), y + Inches(0.15), Inches(2.5), Inches(0.55), fill_color=BG_CARD_ALT)
    add_text(slide, right_text, Inches(8.9), y + Inches(0.2), Inches(2.5), Inches(0.45),
             font_size=12, bold=True, color=col, align=PP_ALIGN.CENTER, font_name="Consolas")

    # Arrow between boxes
    if i < len(layers) - 1:
        arrow_y = y + box_h
        add_text(slide, "^", center_x + Inches(1.55), arrow_y - Pt(2), Inches(0.4), Inches(0.35),
                 font_size=18, color=GRAY_DARK, align=PP_ALIGN.CENTER)

add_footer(slide)
add_slide_number(slide, 4)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — 14-DIMENSIONAL FEATURE VECTOR
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 4, "14-Dimensional Wallet Feature Vector", "x_v in R^14 — Four categories of engineered features")

categories = [
    ("Transaction Velocity", ["incoming_tx", "outgoing_tx", "active_days"], CYAN, "3 features"),
    ("Monetary & Gas", ["avg_tx_amount", "max_tx_amount", "balance", "gas_usage"], GREEN, "4 features"),
    ("Graph Centrality", ["neighbor_count", "degree_centrality", "betweenness", "pagerank", "clustering_coeff"], PURPLE, "5 features"),
    ("Diversity & Cross-Chain", ["token_diversity", "cross_chain_tx_count"], ORANGE, "2 features"),
]

card_w = Inches(2.8)
card_h = Inches(4.5)
start_x = Inches(0.7)
gap = Inches(0.25)

for i, (cat_name, features, col, count_text) in enumerate(categories):
    x = start_x + i * (card_w + gap)
    y = Inches(2.0)
    card = add_shape(slide, x, y, card_w, card_h, fill_color=BG_CARD, line_color=col, line_width=Pt(1.5))
    # Header
    add_rect(slide, x + Pt(1), y + Pt(1), card_w - Pt(2), Inches(0.6), fill_color=col)
    add_text(slide, cat_name, x, y + Inches(0.1), card_w, Inches(0.4),
             font_size=14, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER)
    # Count badge
    add_text(slide, count_text, x, y + Inches(0.7), card_w, Inches(0.3),
             font_size=10, color=col, align=PP_ALIGN.CENTER, font_name="Consolas")
    # Features list
    fy = y + Inches(1.1)
    for feat in features:
        add_text(slide, f"  {feat}", x + Inches(0.15), fy, card_w - Inches(0.3), Inches(0.3),
                 font_size=12, color=GRAY_LIGHT, font_name="Consolas")
        fy += Inches(0.45)

# Total label
add_shape(slide, Inches(4.5), Inches(6.65), Inches(4.3), Inches(0.45), fill_color=BG_CARD, line_color=CYAN, line_width=Pt(1))
add_text(slide, "Total: 14 Engineered Features per Wallet Node", Inches(4.5), Inches(6.68), Inches(4.3), Inches(0.4),
         font_size=12, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

add_footer(slide)
add_slide_number(slide, 5)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — T-EGNN MODEL ARCHITECTURE
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 5, "T-EGNN Model Architecture", "Temporal Explainable Graph Neural Network — dual-head design")

# Left: Architecture flow
flow_steps = [
    ("Input: x_v in R^14", "14-dim wallet feature vector", GRAY),
    ("Linear Projection", "LayerNorm(W_0 * x)", CYAN),
    ("3x GCN Layers", "Symmetric normalized message passing", BLUE),
    ("Temporal Self-Attention", "Multi-head scaled dot-product", GREEN),
]

fw = Inches(5.0)
fh = Inches(0.75)
fx = Inches(0.7)
fy_start = Inches(2.0)

for i, (title, desc, col) in enumerate(flow_steps):
    fy = fy_start + i * Inches(0.95)
    box = add_shape(slide, fx, fy, fw, fh, fill_color=BG_CARD, line_color=col, line_width=Pt(2))
    add_text(slide, title, fx + Inches(0.2), fy + Inches(0.08), fw - Inches(0.4), Inches(0.3),
             font_size=14, bold=True, color=col)
    add_text(slide, desc, fx + Inches(0.2), fy + Inches(0.4), fw - Inches(0.4), Inches(0.3),
             font_size=11, color=GRAY_LIGHT, font_name="Consolas")
    # Arrow
    if i < len(flow_steps) - 1:
        add_text(slide, "v", fx + Inches(2.3), fy + fh - Pt(4), Inches(0.4), Inches(0.35),
                 font_size=18, color=GRAY_DARK, align=PP_ALIGN.CENTER)

# Dual output heads
head_y = fy_start + 4 * Inches(0.95)

# Fraud head
fraud_box = add_shape(slide, Inches(0.7), head_y, Inches(2.3), Inches(1.0), fill_color=BG_CARD, line_color=RED, line_width=Pt(2))
add_text(slide, "Fraud Score Head", Inches(0.7), head_y + Inches(0.1), Inches(2.3), Inches(0.3),
         font_size=13, bold=True, color=RED, align=PP_ALIGN.CENTER)
add_text(slide, "Sigmoid(MLP(Z))", Inches(0.7), head_y + Inches(0.4), Inches(2.3), Inches(0.25),
         font_size=10, color=GRAY_LIGHT, align=PP_ALIGN.CENTER, font_name="Consolas")
add_text(slide, "[0.0 - 1.0]", Inches(0.7), head_y + Inches(0.65), Inches(2.3), Inches(0.25),
         font_size=10, color=GRAY, align=PP_ALIGN.CENTER, font_name="Consolas")

# Attribution head
attr_box = add_shape(slide, Inches(3.3), head_y, Inches(2.4), Inches(1.0), fill_color=BG_CARD, line_color=PURPLE, line_width=Pt(2))
add_text(slide, "Attribution Head", Inches(3.3), head_y + Inches(0.1), Inches(2.4), Inches(0.3),
         font_size=13, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
add_text(slide, "Softmax(MLP(Z))", Inches(3.3), head_y + Inches(0.4), Inches(2.4), Inches(0.25),
         font_size=10, color=GRAY_LIGHT, align=PP_ALIGN.CENTER, font_name="Consolas")
add_text(slide, "8 Classes", Inches(3.3), head_y + Inches(0.65), Inches(2.4), Inches(0.25),
         font_size=10, color=GRAY, align=PP_ALIGN.CENTER, font_name="Consolas")

# Right: Key details
rx = Inches(6.5)
ry = Inches(2.0)
detail_card = add_shape(slide, rx, ry, Inches(6.2), Inches(5.0), fill_color=BG_CARD)

details = [
    ("Model Specifications", "", CYAN, True, 16),
    ("Layers: 3x GCN + 1x Temporal Attention", "", GRAY_LIGHT, False, 12),
    ("Hidden Dim: 128 -> 64 -> 32", "", GRAY_LIGHT, False, 12),
    ("Attention Heads: 4", "", GRAY_LIGHT, False, 12),
    ("Activation: LeakyReLU (a=0.2)", "", GRAY_LIGHT, False, 12),
    ("", "", GRAY_LIGHT, False, 8),
    ("Multi-Task Loss Function", "", GREEN, True, 16),
    ("L_total = 1.0 * BCE + 0.5 * CE + 1e-5 * L2", "", GRAY_LIGHT, False, 12),
    ("", "", GRAY_LIGHT, False, 8),
    ("Wallet Categories (8 Classes)", "", PURPLE, True, 16),
    ("Unknown | Exchange | DeFi | NFT", "", GRAY_LIGHT, False, 12),
    ("Gaming | Mixer | Merchant | Scam", "", GRAY_LIGHT, False, 12),
    ("", "", GRAY_LIGHT, False, 8),
    ("Key Innovations", "", ORANGE, True, 16),
    ("+ Residual skip connections", "", GRAY_LIGHT, False, 12),
    ("+ LayerNorm at each GCN layer", "", GRAY_LIGHT, False, 12),
    ("+ Temporal attention over node embeddings", "", GRAY_LIGHT, False, 12),
]

dy = ry + Inches(0.2)
for text, _, col, bold, fs in details:
    if text:
        add_text(slide, text, rx + Inches(0.3), dy, Inches(5.6), Inches(0.3),
                 font_size=fs, bold=bold, color=col, font_name="Consolas" if not bold else "Calibri")
    dy += Inches(0.25)

add_footer(slide)
add_slide_number(slide, 6)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — MATHEMATICAL FORMULATION
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 6, "Mathematical Formulation", "Core equations driving the T-EGNN model")

# Three formula cards
formulas = [
    ("1. Spatial GCN Message Passing",
     "H^(l+1) = sigma(D^(-1/2) * A_tilde * D^(-1/2) * H^(l) * W^(l))",
     "Symmetric normalization with self-loops; residual skip + LayerNorm",
     CYAN),
    ("2. Multi-Head Temporal Attention",
     "A_temp = Softmax(Q * K^T / sqrt(d_k)) * V",
     "Scaled dot-product attention captures temporal transaction dynamics",
     GREEN),
    ("3. Multi-Task Objective",
     "L = 1.0*BCE + 0.5*CE + 1e-5*||Theta||_2^2",
     "Joint optimization: fraud scoring + wallet attribution + L2 regularization",
     PURPLE),
]

fw = Inches(3.7)
fh = Inches(4.5)
start_x = Inches(0.7)
gap = Inches(0.35)

for i, (title, eq, desc, col) in enumerate(formulas):
    x = start_x + i * (fw + gap)
    y = Inches(2.0)
    card = add_shape(slide, x, y, fw, fh, fill_color=BG_CARD, line_color=col, line_width=Pt(2))

    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), y + Inches(0.15), Inches(0.4), Inches(0.4))
    circle.fill.solid()
    circle.fill.fore_color.rgb = col
    circle.line.fill.background()
    add_text(slide, str(i+1), x + Inches(0.15), y + Inches(0.17), Inches(0.4), Inches(0.35),
             font_size=16, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER, font_name="Consolas")

    add_text(slide, title.split(". ", 1)[1], x + Inches(0.65), y + Inches(0.2), fw - Inches(0.8), Inches(0.35),
             font_size=14, bold=True, color=col)

    # Equation box
    eq_box = add_shape(slide, x + Inches(0.15), y + Inches(0.7), fw - Inches(0.3), Inches(1.6),
                       fill_color=BG_CARD_ALT, line_color=RGBColor(0x33, 0x40, 0x55))
    add_text(slide, eq, x + Inches(0.25), y + Inches(0.85), fw - Inches(0.5), Inches(1.2),
             font_size=13, color=CYAN, font_name="Consolas", line_spacing=1.4)

    # Description
    add_text(slide, desc, x + Inches(0.15), y + Inches(2.5), fw - Inches(0.3), Inches(1.8),
             font_size=12, color=GRAY_LIGHT, line_spacing=1.4)

add_footer(slide)
add_slide_number(slide, 7)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — EXPERIMENTAL RESULTS
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 7, "Experimental Benchmark & Results", "250,000 wallets | 80/10/10 stratified temporal split")

# Table header
table_data = [
    ["Model", "Accuracy", "Precision", "Recall", "F1-Score", "AUC-ROC", "Latency"],
    ["Random Forest", "84.1%", "0.812", "0.795", "0.803", "0.862", "1.2 ms"],
    ["XGBoost", "87.6%", "0.854", "0.831", "0.842", "0.895", "2.5 ms"],
    ["Standard GCN", "89.8%", "0.881", "0.875", "0.878", "0.921", "8.4 ms"],
    ["GAT", "91.5%", "0.902", "0.891", "0.896", "0.942", "14.8 ms"],
    ["T-EGNN (Ours)", "94.2%", "0.938", "0.894", "0.915", "0.963", "11.6 ms"],
]

rows = len(table_data)
cols = len(table_data[0])
tbl_x = Inches(0.7)
tbl_y = Inches(2.0)
tbl_w = Inches(7.5)
tbl_h = Inches(3.0)

table_shape = slide.shapes.add_table(rows, cols, tbl_x, tbl_y, tbl_w, tbl_h)
table = table_shape.table

col_widths = [Inches(1.6), Inches(1.0), Inches(1.0), Inches(0.9), Inches(1.0), Inches(1.0), Inches(1.0)]
for i, w in enumerate(col_widths):
    table.columns[i].width = w

for r in range(rows):
    for c in range(cols):
        cell = table.cell(r, c)
        cell.text = table_data[r][c]
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.name = "Consolas"
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
            if r == 0:  # Header
                p.font.bold = True
                p.font.color.rgb = BG_DARK
                cell.fill.solid()
                cell.fill.fore_color.rgb = CYAN
            elif r == rows - 1:  # Ours - highlighted
                p.font.bold = True
                p.font.color.rgb = GREEN
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x0A, 0x2E, 0x1E)
            else:
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = BG_CARD if r % 2 == 1 else BG_CARD_ALT

# Key observations — right side
obs_x = Inches(8.5)
obs_y = Inches(2.0)
obs_card = add_shape(slide, obs_x, obs_y, Inches(4.2), Inches(3.0), fill_color=BG_CARD)

add_text(slide, "Key Observations", obs_x + Inches(0.2), obs_y + Inches(0.15), Inches(3.8), Inches(0.35),
         font_size=15, bold=True, color=CYAN)

observations = [
    ("+10.1%", "accuracy over Random Forest baseline", GREEN),
    ("+4.4%", "from adding Temporal Attention to GCN", CYAN),
    ("11.6ms", "real-time inference (< 50ms requirement)", ORANGE),
    ("0.963", "AUC-ROC — best among all models", PURPLE),
]

oy = obs_y + Inches(0.6)
for metric, desc, col in observations:
    add_text(slide, metric, obs_x + Inches(0.3), oy, Inches(1.0), Inches(0.3),
             font_size=14, bold=True, color=col, font_name="Consolas")
    add_text(slide, desc, obs_x + Inches(1.35), oy, Inches(2.6), Inches(0.3),
             font_size=11, color=GRAY_LIGHT)
    oy += Inches(0.55)

# Bar chart - AUC-ROC comparison
chart_data = CategoryChartData()
chart_data.categories = ['RF', 'XGBoost', 'GCN', 'GAT', 'T-EGNN']
chart_data.add_series('AUC-ROC', (0.862, 0.895, 0.921, 0.942, 0.963))

chart_frame = slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(0.7), Inches(5.2), Inches(5.5), Inches(1.7),
    chart_data
)
chart = chart_frame.chart
chart.has_legend = False
chart.has_title = False

# Style the chart
plot = chart.plots[0]
plot.gap_width = 100
series = plot.series[0]
series.format.fill.solid()
series.format.fill.fore_color.rgb = CYAN

# Value axis
value_axis = chart.value_axis
value_axis.minimum_scale = 0.8
value_axis.maximum_scale = 1.0
value_axis.major_gridlines.format.line.color.rgb = RGBColor(0x33, 0x40, 0x55)
value_axis.format.line.color.rgb = GRAY_DARK
value_axis.tick_labels.font.size = Pt(9)
value_axis.tick_labels.font.color.rgb = GRAY
value_axis.tick_labels.font.name = "Consolas"

# Category axis
cat_axis = chart.category_axis
cat_axis.format.line.color.rgb = GRAY_DARK
cat_axis.tick_labels.font.size = Pt(9)
cat_axis.tick_labels.font.color.rgb = GRAY
cat_axis.tick_labels.font.name = "Consolas"

# Data labels
plot.has_data_labels = True
data_labels = plot.data_labels
data_labels.font.size = Pt(9)
data_labels.font.color.rgb = WHITE
data_labels.font.name = "Consolas"
data_labels.number_format = '0.000'
data_labels.label_position = XL_LABEL_POSITION.OUTSIDE_END

# Chart area
chart.chart_style = 2

add_footer(slide)
add_slide_number(slide, 8)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — EXPLAINABLE AI & FEATURE IMPORTANCE
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 8, "Explainable AI & Feature Importance", "Gradient-based normalized feature attribution for forensic transparency")

# Feature importance chart
chart_data = CategoryChartData()
chart_data.categories = [
    'Degree Centrality',
    'Max Tx Amount',
    'Cross-Chain Count',
    'Clustering Coeff',
    'Active Days',
    'Neighbor Count',
    'Other'
]
chart_data.add_series('Importance %', (24.5, 21.2, 15.8, 13.4, 9.1, 6.8, 9.2))

chart_frame = slide.shapes.add_chart(
    XL_CHART_TYPE.BAR_CLUSTERED,
    Inches(0.7), Inches(2.0), Inches(6.0), Inches(4.2),
    chart_data
)
chart = chart_frame.chart
chart.has_legend = False
chart.has_title = False

plot = chart.plots[0]
plot.gap_width = 80
series = plot.series[0]

# Color each bar differently
colors = [CYAN, GREEN, ORANGE, PURPLE, BLUE, PINK, GRAY_DARK]
for i, color in enumerate(colors):
    pt = series.points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = color

# Axes
value_axis = chart.value_axis
value_axis.maximum_scale = 30
value_axis.major_gridlines.format.line.color.rgb = RGBColor(0x33, 0x40, 0x55)
value_axis.format.line.color.rgb = GRAY_DARK
value_axis.tick_labels.font.size = Pt(9)
value_axis.tick_labels.font.color.rgb = GRAY
value_axis.tick_labels.font.name = "Consolas"

cat_axis = chart.category_axis
cat_axis.format.line.color.rgb = GRAY_DARK
cat_axis.tick_labels.font.size = Pt(10)
cat_axis.tick_labels.font.color.rgb = WHITE
cat_axis.tick_labels.font.name = "Consolas"

plot.has_data_labels = True
data_labels = plot.data_labels
data_labels.font.size = Pt(10)
data_labels.font.color.rgb = WHITE
data_labels.font.name = "Consolas"
data_labels.number_format = '0.0"%"'
data_labels.label_position = XL_LABEL_POSITION.OUTSIDE_END

# Right: Formula + explanation
rx = Inches(7.0)
ry = Inches(2.0)
card = add_shape(slide, rx, ry, Inches(5.8), Inches(4.8), fill_color=BG_CARD)

add_text(slide, "Attribution Formula", rx + Inches(0.3), ry + Inches(0.2), Inches(5.2), Inches(0.3),
         font_size=16, bold=True, color=CYAN)

formula_box = add_shape(slide, rx + Inches(0.3), ry + Inches(0.6), Inches(5.2), Inches(0.8),
                        fill_color=BG_CARD_ALT, line_color=RGBColor(0x33, 0x40, 0x55))
add_text(slide, "S(x_i) = |df/dx_i * x_i| / sum_j|df/dx_j * x_j|",
         rx + Inches(0.5), ry + Inches(0.72), Inches(4.8), Inches(0.5),
         font_size=13, color=CYAN, font_name="Consolas")

insights = [
    ("Degree Centrality (24.5%)", "Spikes in connection count signal rapid fund disbursement — peeling chains create many one-time counterparties"),
    ("Max Tx Amount (21.2%)", "Peak single burst transfers indicate structuring or mixer deposit behavior"),
    ("Cross-Chain Count (15.8%)", "Frequent asset swaps across chain bridges suggest cross-chain evasion tactics"),
    ("Clustering Coefficient (13.4%)", "Interconnected relay rings between wallets in known mixer networks"),
]

iy = ry + Inches(1.6)
for title, desc in insights:
    add_rect(slide, rx + Inches(0.3), iy, Pt(3), Inches(0.7), fill_color=CYAN)
    add_text(slide, title, rx + Inches(0.55), iy, Inches(5.0), Inches(0.3),
             font_size=12, bold=True, color=WHITE)
    add_text(slide, desc, rx + Inches(0.55), iy + Inches(0.3), Inches(5.0), Inches(0.45),
             font_size=10, color=GRAY_LIGHT, line_spacing=1.3)
    iy += Inches(0.78)

add_footer(slide)
add_slide_number(slide, 9)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — FORENSIC CASE STUDY
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 9, "Forensic Case Study: Mixer Obfuscation", "Real-world detection scenario with automated reasoning")

# Left: Target wallet info
lx = Inches(0.7)
ly = Inches(2.0)

# Target card
target_card = add_shape(slide, lx, ly, Inches(6.0), Inches(1.5), fill_color=BG_CARD, line_color=RED, line_width=Pt(2))
add_text(slide, "TARGET ADDRESS", lx + Inches(0.3), ly + Inches(0.15), Inches(5.4), Inches(0.3),
         font_size=11, bold=True, color=RED, font_name="Consolas")
add_text(slide, "0x742d35Cc6634C0532925a3b844Bc454e4438f44e",
         lx + Inches(0.3), ly + Inches(0.5), Inches(5.4), Inches(0.3),
         font_size=12, color=GRAY_LIGHT, font_name="Consolas")

# Risk metrics
risk_y = ly + Inches(0.9)
for i, (label, value, col) in enumerate([
    ("FRAUD SCORE", "92.4%", RED),
    ("CATEGORY", "Mixer", PURPLE),
    ("CONFIDENCE", "96.1%", GREEN),
]):
    bx = lx + Inches(0.3) + i * Inches(1.8)
    add_text(slide, label, bx, risk_y, Inches(1.6), Inches(0.2),
             font_size=9, bold=True, color=GRAY, font_name="Consolas")
    add_text(slide, value, bx, risk_y + Inches(0.2), Inches(1.6), Inches(0.3),
             font_size=18, bold=True, color=col, font_name="Consolas")

# Sub-graph visualization
sg_y = Inches(3.8)
sg_card = add_shape(slide, lx, sg_y, Inches(6.0), Inches(3.0), fill_color=BG_CARD)

add_text(slide, "LOCAL TRANSACTION SUB-GRAPH", lx + Inches(0.3), sg_y + Inches(0.1), Inches(5.4), Inches(0.3),
         font_size=11, bold=True, color=CYAN, font_name="Consolas")

# Node representations
nodes = [
    ("Victim Wallet", "Risk: 12.0%", Inches(1.2), Inches(4.5), GREEN),
    ("Suspect (Target)", "Risk: 92.4%", Inches(3.4), Inches(4.3), RED),
    ("Mixer Relay #1", "Risk: 98.1%", Inches(1.5), Inches(5.7), RED),
    ("Mixer Relay #2", "Risk: 95.7%", Inches(3.4), Inches(5.7), ORANGE),
    ("Mixer Relay #3", "Risk: 99.0%", Inches(5.3), Inches(5.7), RED),
]

for name, risk, nx, ny, col in nodes:
    node = add_shape(slide, nx, ny, Inches(1.4), Inches(0.8), fill_color=BG_CARD_ALT, line_color=col, line_width=Pt(1.5))
    add_text(slide, name, nx, ny + Inches(0.08), Inches(1.4), Inches(0.3),
             font_size=9, bold=True, color=col, align=PP_ALIGN.CENTER, font_name="Consolas")
    add_text(slide, risk, nx, ny + Inches(0.4), Inches(1.4), Inches(0.25),
             font_size=8, color=GRAY_LIGHT, align=PP_ALIGN.CENTER, font_name="Consolas")

# Right: Automated reasoning
rx = Inches(7.0)
ry = Inches(2.0)
reason_card = add_shape(slide, rx, ry, Inches(5.8), Inches(4.8), fill_color=BG_CARD)

add_text(slide, "Automated Reasoning Factors", rx + Inches(0.3), ry + Inches(0.15), Inches(5.2), Inches(0.35),
         font_size=16, bold=True, color=CYAN)

reasons = [
    ("Risk Factor 1", "Abnormally high degree centrality coupled with low active days — signature of a rapid peeling chain pattern", RED),
    ("Risk Factor 2", "Outbound transaction split evenly into identical micro-amounts across 85 distinct counterparties within a 12-minute window", ORANGE),
    ("Risk Factor 3", "High clustering coefficient (0.78) with verified blacklisted mixer relay contracts on-chain", PURPLE),
]

ry_pos = ry + Inches(0.65)
for title, desc, col in reasons:
    # Numbered circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, rx + Inches(0.3), ry_pos + Inches(0.05), Inches(0.35), Inches(0.35))
    circle.fill.solid()
    circle.fill.fore_color.rgb = col
    circle.line.fill.background()
    idx = reasons.index((title, desc, col)) + 1
    add_text(slide, str(idx), rx + Inches(0.3), ry_pos + Inches(0.07), Inches(0.35), Inches(0.3),
             font_size=12, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER, font_name="Consolas")

    add_text(slide, title, rx + Inches(0.8), ry_pos, Inches(4.7), Inches(0.3),
             font_size=13, bold=True, color=col)
    add_text(slide, desc, rx + Inches(0.8), ry_pos + Inches(0.35), Inches(4.7), Inches(0.7),
             font_size=11, color=GRAY_LIGHT, line_spacing=1.3)
    ry_pos += Inches(1.15)

add_footer(slide)
add_slide_number(slide, 10)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 11 — FULL-STACK IMPLEMENTATION
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 10, "Full-Stack Production Implementation", "Enterprise-grade platform with real-time capabilities")

# Tech stack columns
stacks = [
    ("Backend", CYAN, [
        "FastAPI (Python 3.11)",
        "SQLAlchemy ORM",
        "Neo4j Driver",
        "PyTorch Geometric",
        "JWT Authentication",
        "14 API Routers",
    ]),
    ("Frontend", GREEN, [
        "React 18 + Vite",
        "Tailwind CSS",
        "Cytoscape.js",
        "Recharts",
        "Zustand State Mgmt",
        "13 Pages",
    ]),
    ("Infrastructure", PURPLE, [
        "Docker Compose",
        "PostgreSQL 16",
        "Neo4j 5.x",
        "Alembic Migrations",
        "JWT Security",
        "Role-Based Access",
    ]),
]

col_w = Inches(3.7)
col_h = Inches(4.2)
start_x = Inches(0.7)
gap = Inches(0.35)

for i, (title, col, items) in enumerate(stacks):
    x = start_x + i * (col_w + gap)
    y = Inches(2.0)
    card = add_shape(slide, x, y, col_w, col_h, fill_color=BG_CARD, line_color=col, line_width=Pt(2))

    # Title bar
    add_rect(slide, x + Pt(1), y + Pt(1), col_w - Pt(2), Inches(0.5), fill_color=col)
    add_text(slide, title, x, y + Inches(0.08), col_w, Inches(0.35),
             font_size=15, bold=True, color=BG_DARK, align=PP_ALIGN.CENTER)

    iy = y + Inches(0.7)
    for item in items:
        add_text(slide, f"  {item}", x + Inches(0.2), iy, col_w - Inches(0.4), Inches(0.35),
                 font_size=12, color=GRAY_LIGHT, font_name="Consolas")
        iy += Inches(0.5)

# Key modules section
modules_y = Inches(6.4)
add_text(slide, "Key Application Modules:", Inches(0.7), modules_y, Inches(2.5), Inches(0.3),
         font_size=12, bold=True, color=CYAN)

modules = [
    "Fraud Detection", "Wallet Attribution", "Graph Canvas",
    "Investigation Cases", "Blacklist CRUD", "PDF Reports",
]
mx = Inches(3.3)
for mod in modules:
    badge = add_shape(slide, mx, modules_y - Inches(0.02), Inches(1.5), Inches(0.35), fill_color=BG_CARD_ALT, line_color=CYAN_DIM)
    add_text(slide, mod, mx, modules_y + Inches(0.02), Inches(1.5), Inches(0.3),
             font_size=9, color=CYAN_DIM, align=PP_ALIGN.CENTER, font_name="Consolas")
    mx += Inches(1.65)

add_footer(slide)
add_slide_number(slide, 11)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 12 — CONCLUSION & FUTURE WORK
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)
add_section_header(slide, 11, "Conclusion & Future Roadmap", "Summary of achievements and directions for future research")

# Left: Achievements
ach_x = Inches(0.7)
ach_y = Inches(2.0)
ach_card = add_shape(slide, ach_x, ach_y, Inches(5.8), Inches(4.8), fill_color=BG_CARD, line_color=GREEN, line_width=Pt(2))

add_text(slide, "Key Achievements", ach_x + Inches(0.3), ach_y + Inches(0.2), Inches(5.2), Inches(0.4),
         font_size=18, bold=True, color=GREEN)

achievements = [
    ("Unified 5-Chain Graph Model", "BTC, ETH, BNB, MATIC, TRX in a single integrated framework"),
    ("94.2% Classification Accuracy", "Outperforming all baselines including GAT and XGBoost"),
    ("0.963 AUC-ROC Score", "State-of-the-art discrimination between fraudulent and legitimate wallets"),
    ("11.6ms Inference Latency", "Sub-12ms real-time query speed for exchange compliance"),
    ("Native Explainability", "Transparent feature attributions without post-hoc overhead"),
]

ay = ach_y + Inches(0.7)
for title, desc in achievements:
    add_text(slide, "v", ach_x + Inches(0.3), ay, Inches(0.3), Inches(0.25),
             font_size=14, bold=True, color=GREEN, font_name="Consolas")
    add_text(slide, title, ach_x + Inches(0.65), ay, Inches(4.9), Inches(0.25),
             font_size=13, bold=True, color=WHITE)
    add_text(slide, desc, ach_x + Inches(0.65), ay + Inches(0.3), Inches(4.9), Inches(0.3),
             font_size=11, color=GRAY_LIGHT)
    ay += Inches(0.7)

# Right: Future Work
fut_x = Inches(6.8)
fut_y = Inches(2.0)
fut_card = add_shape(slide, fut_x, fut_y, Inches(5.8), Inches(4.8), fill_color=BG_CARD, line_color=PURPLE, line_width=Pt(2))

add_text(slide, "Future Roadmap", fut_x + Inches(0.3), fut_y + Inches(0.2), Inches(5.2), Inches(0.4),
         font_size=18, bold=True, color=PURPLE)

futures = [
    ("Zero-Knowledge Fraud Verification", "ZK-SNARK auditing without revealing user balance privacy; regulatory compliance with privacy", CYAN),
    ("Streaming Dynamic GNNs", "Real-time block stream ingestion via PyTorch Geometric Temporal for live transaction monitoring", GREEN),
    ("L2 Network Expansion", "Arbitrum, Optimism, ZK-Sync, and Solana network support for comprehensive coverage", ORANGE),
]

fy = fut_y + Inches(0.7)
for title, desc, col in futures:
    # Arrow indicator
    add_rect(slide, fut_x + Inches(0.3), fy, Pt(4), Inches(0.9), fill_color=col)
    add_text(slide, title, fut_x + Inches(0.6), fy, Inches(5.0), Inches(0.3),
             font_size=13, bold=True, color=col)
    add_text(slide, desc, fut_x + Inches(0.6), fy + Inches(0.35), Inches(5.0), Inches(0.55),
             font_size=11, color=GRAY_LIGHT, line_spacing=1.3)
    fy += Inches(1.15)

add_footer(slide)
add_slide_number(slide, 12)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 13 — THANK YOU
# ══════════════════════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(BLANK_LAYOUT)
set_slide_bg(slide)

# Top accent bar
add_rect(slide, Inches(0), Inches(0), SLIDE_W, Pt(4), fill_color=CYAN)

# Shield icon
shield = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(5.9), Inches(1.5), Inches(1.5), Inches(1.5))
shield.fill.solid()
shield.fill.fore_color.rgb = RGBColor(0x00, 0x3D, 0x45)
shield.line.color.rgb = CYAN
shield.line.width = Pt(2)
add_text(slide, "AI", Inches(5.9), Inches(1.8), Inches(1.5), Inches(0.8),
         font_size=28, bold=True, color=CYAN, align=PP_ALIGN.CENTER, font_name="Consolas")

add_text(slide, "Thank You!", Inches(2), Inches(3.3), Inches(9.3), Inches(0.8),
         font_size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

add_text(slide, "Questions & Discussion", Inches(2), Inches(4.1), Inches(9.3), Inches(0.5),
         font_size=20, color=GRAY_LIGHT, align=PP_ALIGN.CENTER)

add_accent_line(slide, Inches(5.2), Inches(4.7), Inches(3), CYAN)

# Contact info
contact_items = [
    ("Author:", "Nitin S.", CYAN),
    ("Email:", "nitin@cryptoshield.ai", GREEN),
    ("Repository:", "Nitins2005/blockchain", PURPLE),
]

cy = Inches(5.1)
for label, value, col in contact_items:
    add_text(slide, label, Inches(4.5), cy, Inches(1.5), Inches(0.3),
             font_size=13, bold=True, color=GRAY, align=PP_ALIGN.RIGHT)
    add_text(slide, value, Inches(6.1), cy, Inches(3.5), Inches(0.3),
             font_size=13, color=col, font_name="Consolas")
    cy += Inches(0.35)

# Bottom stats recap
metrics = [
    ("94.2%", "Accuracy", GREEN),
    ("0.963", "AUC-ROC", CYAN),
    ("11.6ms", "Latency", PURPLE),
    ("5 Chains", "Multi-Chain", ORANGE),
]
card_w = Inches(2.1)
start_x = Inches(1.8)
gap = Inches(0.3)
for i, (val, label, col) in enumerate(metrics):
    x = start_x + i * (card_w + gap)
    card = add_shape(slide, x, Inches(6.2), card_w, Inches(0.7), fill_color=BG_CARD, line_color=col, line_width=Pt(1))
    add_text(slide, f"{val}  {label}", x, Inches(6.28), card_w, Inches(0.5),
             font_size=12, bold=True, color=col, align=PP_ALIGN.CENTER, font_name="Consolas")

add_slide_number(slide, 13)


# ─── Save ────────────────────────────────────────────────────────────────────
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "CryptoShield_AI_Presentation.pptx")
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
